# -*- coding: utf-8 -*-
import docx

for name in ['初稿1', '初稿2']:
    path = rf'C:\Users\ZhuanZ（无密码）\Desktop\ffffffffff\毕业论文\{name}.docx'
    try:
        d = docx.Document(path)
    except Exception as e:
        print(f'{name}: 打开失败 {e}')
        continue
    print(f'========== {name}.docx ==========')
    print(f'段落数: {len(d.paragraphs)}, 表格数: {len(d.tables)}')
    print('--- 段落内容 ---')
    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        if t:
            print(f'[{i}] {t}')
    print()
    for ti, tb in enumerate(d.tables):
        print(f'--- 表格{ti+1} ({len(tb.rows)}x{len(tb.columns)}) ---')
        for r in tb.rows:
            print(' | '.join(c.text.replace(chr(10),' ')[:40] for c in r.cells))
    print()
