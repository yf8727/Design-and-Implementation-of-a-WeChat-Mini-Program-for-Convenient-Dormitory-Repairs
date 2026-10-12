# -*- coding: utf-8 -*-
import docx
from lxml import etree

path = r'C:\Users\ZhuanZ（无密码）\Desktop\ffffffffff\毕业论文\初稿1.docx'
d = docx.Document(path)
body = d.element.body
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

print('=== 封面表格(表格1)结构 ===')
t = d.tables[0]
for ri, row in enumerate(t.rows):
    for ci, cell in enumerate(row.cells):
        print(f'  行{ri}列{ci}: "{cell.text}"')

print()
print('=== 评阅表(表格4)结构 ===')
t4 = d.tables[3]
for ri, row in enumerate(t4.rows):
    seen = set()
    for ci, cell in enumerate(row.cells):
        tc_id = id(cell._tc)
        if tc_id in seen:
            continue
        seen.add(tc_id)
        print(f'  行{ri}单元格[{ci}]: "{cell.text[:50]}"')

print()
print('=== 目录SDT结构 ===')
for ch in body:
    if etree.QName(ch).localname == 'sdt':
        sdt = ch
        sdtContent = None
        for sc in sdt:
            if etree.QName(sc).localname == 'sdtContent':
                sdtContent = sc
                break
        if sdtContent is not None:
            n = len(list(sdtContent))
            print(f'  sdtContent子元素数: {n}')
            for i, child in enumerate(list(sdtContent)):
                texts = ''.join(t.text or '' for t in child.iter(W + 't'))
                print(f'    [{i}] {etree.QName(child).localname}: "{texts[:40]}"')
        break

print()
print('=== 分节(sectPr)位置 ===')
for i, ch in enumerate(list(body)):
    if etree.QName(ch).localname == 'p':
        sp = ch.find(W + 'pPr/' + W + 'sectPr')
        if sp is not None:
            texts = ''.join(t.text or '' for t in ch.iter(W + 't'))
            print(f'  body[{i}] 含sectPr 文本="{texts[:20]}"')

print()
print('=== 需要修改的段落定位 ===')
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if any(k in t for k in ['两类用户', '陈淼等针对传统', '国外高校在设施', '综合分析国内外', '触手可及', '四种文件组成', '云开发是腾讯云', '贾志杰', '参考文献', '围绕宿舍报修业务流程']):
        print(f'  段落[{i}]: "{t[:60]}"')

print()
print('=== 5.3结束、6章开始的段落 ===')
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t.startswith('（3）报修数据统计') or t == '6 系统测试' or t.startswith('5.3') or t.startswith('6.1'):
        print(f'  段落[{i}]: "{t[:60]}"')
