# -*- coding: utf-8 -*-
import docx
from lxml import etree

path = r'C:\Users\ZhuanZ（无密码）\Desktop\ffffffffff\毕业论文\work_tmp\SOURCE.docx'
doc = docx.Document(path)
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

print('=== 检查节2/3/4页脚是否为PAGE域 ===')
for i, sec in enumerate(doc.sections):
    f = sec.footer
    xml = f._element.xml
    has_page = 'PAGE' in xml
    has_fldChar = 'fldChar' in xml
    txts = [p.text for p in f.paragraphs]
    print(f'节{i}: 页脚文本={txts}, 含PAGE域字串={has_page}, 含fldChar={has_fldChar}')

print()
print('=== 检查说明文本框(txbxContent)位置 ===')
body = doc.element.body
for i, p in enumerate(doc.paragraphs):
    xml = p._element.xml
    if 'txbxContent' in xml:
        # 提取文本框中的文字
        texts = []
        for t in p._element.iter(W + 't'):
            texts.append(t.text or '')
        joined = ''.join(texts)
        print(f'段落[{i}] 含文本框: "{joined[:60]}"')

print()
print('=== 检查封面题目行的占位符 ===')
p = doc.paragraphs[14]
print('段落[14] runs:', [r.text for r in p.runs])
print('段落[14] XML片段:')
print(p._element.xml[:2000])
