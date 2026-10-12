# -*- coding: utf-8 -*-
import zipfile, re
path = r'C:\Users\ZhuanZ（无密码）\Desktop\ffffffffff\毕业论文\专科毕业论文（设计）模板及要求(1).docx'
z = zipfile.ZipFile(path)

xml = z.read('word/document.xml').decode('utf-8')
print("===== document.xml 中所有 sectPr =====")
for m in re.finditer(r'<w:sectPr[^>]*>.*?</w:sectPr>|<w:sectPr[^>]*/>', xml, re.S):
    s = m.group(0)
    refs = re.findall(r'<w:(footer|header)Reference w:type="([^"]*)" r:id="(rId\d+)"', s)
    titlepg = 'titlePg' in s
    pg = re.findall(r'<w:pgNumType[^>]*/>', s)
    print(f'  refs={refs} titlePg={titlepg} pgNumType={pg}')

print()
print("===== rels =====")
rels = z.read('word/_rels/document.xml.rels').decode('utf-8')
for rid, target in re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels):
    print(f'  {rid} -> {target}')

print()
print("===== 检查每个节实际的footer引用（从document.xml上下文看） =====")
# 按顺序找 sectPr 及其之前的 footerReference
parts = re.split(r'(<w:sectPr[^>]*>.*?</w:sectPr>|<w:sectPr[^>]*/>)', xml, flags=re.S)
sec = 0
for part in parts:
    if part.startswith('<w:sectPr'):
        refs = re.findall(r'<w:(footer|header)Reference w:type="([^"]*)" r:id="(rId\d+)"', part)
        titlepg = 'titlePg' in part
        print(f'  节{sec}: refs={refs} titlePg={titlepg}')
        sec += 1

print()
print("===== settings.xml =====")
settings = z.read('word/settings.xml').decode('utf-8')
print('  evenAndOddHeaders' in settings, 'defaultTabStop' in settings)

print()
print("===== 封面表格详细结构（第一个表格） =====")
import docx
from docx.oxml.ns import qn
doc = docx.Document(path)
tbl = doc.tables[0]
for r_i, row in enumerate(tbl.rows):
    print(f'  Row {r_i}: 高度={row.height}')
    for c_i, cell in enumerate(row.cells):
        # 单元格宽度
        tcW = cell._tc.find(qn('w:tcPr'))
        w = None
        if tcW is not None:
            tcW_e = tcW.find(qn('w:tcW'))
            if tcW_e is not None:
                w = tcW_e.get(qn('w:w'))
        print(f'    Cell[{c_i}]: w={w} text={cell.text!r}')

print()
print("===== 评阅表结构（第二个表格） =====")
tbl2 = doc.tables[1]
for r_i, row in enumerate(tbl2.rows):
    seen = set()
    cells = []
    for c in row.cells:
        if id(c._tc) in seen: continue
        seen.add(id(c._tc))
        cells.append(c.text.replace('\n','|'))
    print(f'  Row {r_i}: {cells}')
