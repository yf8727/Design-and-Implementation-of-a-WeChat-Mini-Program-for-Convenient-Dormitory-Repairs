# -*- coding: utf-8 -*-
"""
构建论文：宿舍便民报修微信小程序的设计与实现（完整版）
基于江西软件职业技术大学专科毕业论文模板
"""
import shutil, os
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.text.paragraph import Paragraph

SRC = r'C:\Users\ZhuanZ（无密码）\Desktop\ffffffffff\毕业论文\专科毕业论文（设计）模板及要求(1).docx'
TGT = r'C:\Users\ZhuanZ（无密码）\Desktop\ffffffffff\毕业论文\宿舍便民报修微信小程序的设计与实现.docx'
FIG = r'C:\Users\ZhuanZ（无密码）\Desktop\ffffffffff\毕业论文\work\figs'

# ---------- 准备目标 ----------
if os.path.exists(TGT):
    try:
        os.chmod(TGT, 0o666); os.remove(TGT)
    except Exception: pass
shutil.copy2(SRC, TGT)
try:
    os.chmod(TGT, 0o666)
except Exception: pass
doc = docx.Document(TGT)

# ================= 工具函数 =================
def run_eastasia(r):
    rPr = r._element.rPr
    if rPr is None or rPr.rFonts is None:
        return None
    return rPr.rFonts.get(qn('w:eastAsia'))

def set_run_font(run, eastasia='宋体', ascii_f='Times New Roman', size_pt=12, bold=None, color='000000'):
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    try:
        run.font.color.rgb = RGBColor.from_string(color)
    except Exception:
        pass
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), ascii_f); rFonts.set(qn('w:hAnsi'), ascii_f)
    rFonts.set(qn('w:eastAsia'), eastasia); rFonts.set(qn('w:cs'), ascii_f)

def set_para(p, align=None, line_exact=None, first_line_chars=None, before=None, after=None,
             left_indent_twips=None, hanging_twips=None, outline=None, page_break_before=False):
    if align is not None:
        p.paragraph_format.alignment = align
    if line_exact is not None:
        p.paragraph_format.line_spacing = Pt(line_exact)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pPr = p._element.get_or_add_pPr()
    if outline is not None:
        ol = OxmlElement('w:outlineLvl'); ol.set(qn('w:val'), str(outline)); pPr.append(ol)
    if page_break_before:
        pb = OxmlElement('w:pageBreakBefore'); pPr.append(pb)
    ind = pPr.find(qn('w:ind'))
    if first_line_chars is not None:
        if ind is None:
            ind = OxmlElement('w:ind'); pPr.append(ind)
        ind.set(qn('w:firstLineChars'), str(int(first_line_chars*100)))
        ind.set(qn('w:firstLine'), str(int(first_line_chars*240)))
    if left_indent_twips is not None:
        if ind is None:
            ind = OxmlElement('w:ind'); pPr.append(ind)
        ind.set(qn('w:left'), str(int(left_indent_twips)))
    if hanging_twips is not None:
        if ind is None:
            ind = OxmlElement('w:ind'); pPr.append(ind)
        ind.set(qn('w:hanging'), str(int(hanging_twips)))
        ind.set(qn('w:hangingChars'), str(int(hanging_twips/24)))
    if before is not None or after is not None:
        sp = pPr.find(qn('w:spacing'))
        if sp is None:
            sp = OxmlElement('w:spacing'); pPr.append(sp)
        if before is not None: sp.set(qn('w:before'), str(int(before*20)))
        if after is not None: sp.set(qn('w:after'), str(int(after*20)))

def new_para(text='', size=12, eastasia='宋体', ascii_f='Times New Roman', bold=None,
             align=None, line_exact=None, first_line_chars=None, before=None, after=None,
             left_indent_twips=None, hanging_twips=None, outline=None, page_break_before=False):
    el = OxmlElement('w:p')
    p = Paragraph(el, doc)
    if text:
        r = p.add_run(text)
        set_run_font(r, eastasia=eastasia, ascii_f=ascii_f, size_pt=size, bold=bold)
    set_para(p, align=align, line_exact=line_exact, first_line_chars=first_line_chars,
             before=before, after=after, left_indent_twips=left_indent_twips,
             hanging_twips=hanging_twips, outline=outline, page_break_before=page_break_before)
    return el

def body_para(text, **kw):
    kw.setdefault('size', 12); kw.setdefault('align', WD_ALIGN_PARAGRAPH.JUSTIFY)
    kw.setdefault('line_exact', 22); kw.setdefault('first_line_chars', 2)
    return new_para(text, **kw)

def h1_para(text):
    return new_para(text, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
                    before=12, after=6, outline=0)

def h2_para(text):
    return new_para(text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT,
                    before=6, after=3, outline=1)

def caption_para(text):
    return new_para(text, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=8)

def image_para(fname, width_cm=13.5):
    el = OxmlElement('w:p')
    p = Paragraph(el, doc)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(os.path.join(FIG, fname), width=Cm(width_cm))
    return el

def code_para(text):
    return new_para(text, size=9, eastasia='宋体', ascii_f='Courier New',
                    align=WD_ALIGN_PARAGRAPH.LEFT, before=2, after=2, left_indent_twips=480)

def insert_after(anchor, elems):
    for e in elems:
        anchor.addnext(e)
        anchor = e

# ================= 1. 封面题目 =================
for p in doc.paragraphs:
    # 封面"题目："标签段落（去空格后为"题目："）
    if p.text.replace(' ', '') == '题目：' and p.text.strip().startswith('题'):
        r = p.add_run('宿舍便民报修微信小程序的设计与实现')
        set_run_font(r, eastasia='宋体', size_pt=15.5, bold=True)
        break

# ================= 2. 样式 =================
styles_el = doc.styles.element
def style_exists(sid):
    return any(s.get(qn('w:styleId')) == sid for s in styles_el.findall(qn('w:style')))

def add_style(styleId, name, outline, size_pt, bold, indent_cm=None):
    st = OxmlElement('w:style')
    st.set(qn('w:type'), 'paragraph'); st.set(qn('w:styleId'), styleId)
    nm = OxmlElement('w:name'); nm.set(qn('w:val'), name); st.append(nm)
    ba = OxmlElement('w:basedOn'); ba.set(qn('w:val'), 'Normal'); st.append(ba)
    qf = OxmlElement('w:qFormat'); st.append(qf)
    pPr = OxmlElement('w:pPr')
    if outline is not None:
        ol = OxmlElement('w:outlineLvl'); ol.set(qn('w:val'), str(outline)); pPr.append(ol)
    kn = OxmlElement('w:keepNext'); pPr.append(kn)
    kl = OxmlElement('w:keepLines'); pPr.append(kl)
    if indent_cm:
        ind = OxmlElement('w:ind'); ind.set(qn('w:left'), str(int(indent_cm*567)))
        ind.set(qn('w:hanging'), str(int(indent_cm*567))); pPr.append(ind)
    st.append(pPr)
    rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'), 'Times New Roman'); rf.set(qn('w:hAnsi'), 'Times New Roman')
    rf.set(qn('w:eastAsia'), '宋体'); rf.set(qn('w:hint'), 'eastAsia'); rPr.append(rf)
    if bold:
        rPr.append(OxmlElement('w:b')); rPr.append(OxmlElement('w:bCs'))
    c = OxmlElement('w:color'); c.set(qn('w:val'), '000000'); rPr.append(c)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(size_pt*2)); rPr.append(sz)
    szcs = OxmlElement('w:szCs'); szcs.set(qn('w:val'), str(size_pt*2)); rPr.append(szcs)
    st.append(rPr)
    return st

for sid, nm, ol, sz, bd in [('Heading1','heading 1',0,14,True), ('Heading2','heading 2',1,12,False)]:
    if not style_exists(sid):
        styles_el.append(add_style(sid, nm, ol, sz, bd))
for sid, nm, sz, bd, icm in [('TOC1','toc 1',14,True,None), ('TOC2','toc 2',12,False,0.75)]:
    if not style_exists(sid):
        styles_el.append(add_style(sid, nm, None, sz, bd, indent_cm=icm))

# ================= 3. settings updateFields =================
settings_el = doc.settings.element
if settings_el.find(qn('w:updateFields')) is None:
    uf = OxmlElement('w:updateFields'); uf.set(qn('w:val'), 'true')
    settings_el.insert(0, uf)

# ================= 4. 目录 =================
toc_p = None
for p in doc.paragraphs:
    if '参考文献' in p.text and '....' in p.text:
        toc_p = p; break
if toc_p is not None:
    for r in list(toc_p.runs):
        r._element.getparent().remove(r._element)
    pPr = toc_p._element.get_or_add_pPr()
    ind = OxmlElement('w:ind'); ind.set(qn('w:firstLine'), '0'); ind.set(qn('w:firstLineChars'), '0'); pPr.append(ind)
    f1 = OxmlElement('w:r'); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'),'begin'); fc.set(qn('w:dirty'),'true'); f1.append(fc)
    f2 = OxmlElement('w:r'); it = OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text = ' TOC \\o "1-2" \\h \\z \\u '; f2.append(it)
    f3 = OxmlElement('w:r'); fc2 = OxmlElement('w:fldChar'); fc2.set(qn('w:fldCharType'),'separate'); f3.append(fc2)
    f4 = OxmlElement('w:r'); rpr = OxmlElement('w:rPr'); sz = OxmlElement('w:sz'); sz.set(qn('w:val'),'24'); rpr.append(sz); t = OxmlElement('w:t'); t.text = '（打开文档后请更新目录）'; f4.append(rpr); f4.append(t)
    f5 = OxmlElement('w:r'); fc3 = OxmlElement('w:fldChar'); fc3.set(qn('w:fldCharType'),'end'); f5.append(fc3)
    for e in (f1,f2,f3,f4,f5):
        toc_p._element.append(e)

# ================= 5. 摘要 =================
abstract_text = ('随着高校办学规模的不断扩大，学生宿舍的设施设备数量持续增加，水电、门窗、家具、网络等故障报修需求日益频繁。'
    '传统的宿舍报修多采用纸质登记或电话报修方式，存在流程繁琐、信息传递不及时、维修进度无法跟踪等问题，影响了学生的正常生活和后勤服务效率。'
    '针对上述问题，本文设计并实现了一个基于微信小程序的宿舍便民报修系统。系统采用微信小程序云开发技术，前端使用微信开发者工具完成页面设计与交互，'
    '后端借助云函数、云数据库和云存储实现业务逻辑处理与数据管理，无需独立搭建和维护服务器。系统面向学生与后勤管理人员两类用户，'
    '实现了在线报修、故障类型分类、报修进度查询、催办、评价反馈，以及报修工单管理、维修人员派单、维修状态更新等主要功能。'
    '测试结果表明，系统各功能模块运行稳定、操作便捷，有效简化了宿舍报修流程，提高了维修响应速度与后勤服务质量，具有一定的实用价值和推广意义。')

# 先删除所有示例楷体段落（两段示例摘要），再插入正式摘要
to_remove = []
for p in list(doc.paragraphs):
    if p.text.strip().startswith('【关键词】'):
        continue
    kaiti_runs = [r for r in p.runs if run_eastasia(r) == '楷体' and r.text.strip()]
    if kaiti_runs and p.text.strip():
        to_remove.append(p)
for p in to_remove:
    p._element.getparent().remove(p._element)

# 在 【关键词】 段之前插入正式摘要（作为独立段）
anchor_kw = None
for p in doc.paragraphs:
    if p.text.strip().startswith('【关键词】'):
        anchor_kw = p; break
if anchor_kw is not None:
    abst_p = new_para(abstract_text, size=12, eastasia='楷体', align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                      line_exact=22, first_line_chars=2)
    anchor_kw._element.addprevious(abst_p)

# 关键词
for p in doc.paragraphs:
    if p.text.strip().startswith('【关键词】'):
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        r1 = p.add_run('【关键词】'); set_run_font(r1, eastasia='黑体', size_pt=12)
        r2 = p.add_run('微信小程序；云开发；宿舍报修；工单管理；云函数'); set_run_font(r2, eastasia='楷体', size_pt=12)
        set_para(p, line_exact=22)
        break

print('阶段1完成：封面/样式/目录/摘要/关键词')

# ================= 6. 删除所有"说明"块 =================
del_patterns = ['说明：', '1.“', '2.“', '3.中', '4.从本页', '5.正文内', '1.“参考', '2.参考文',
                '3.参考文', '4.参考文献数量', '5.参考文献的排序']
# 更精确：删除含"说明："或以数字序号开头的说明段
def is_note_para(p):
    t = p.text.strip()
    if not t:
        return False
    if t == '说明：':
        return True
    if t.startswith('1.') and ('摘要' in t or '目录' in t or '参考' in t):
        return True
    if t.startswith('2.') and ('关键词' in t or '一级标题' in t or '参考' in t):
        return True
    if t.startswith('3.') and ('中文摘要' in t or '参考' in t):
        return True
    if t.startswith('4.') and ('本页' in t or '参考' in t):
        return True
    if t.startswith('5.') and ('正文' in t or '参考' in t):
        return True
    return False

removed_notes = []
for p in list(doc.paragraphs):
    if is_note_para(p):
        removed_notes.append(p)
for p in removed_notes:
    p._element.getparent().remove(p._element)
print(f'删除说明段落 {len(removed_notes)} 个')

# ================= 7. 正文内容 =================
# 找 "正   文" 标签段落
zhengwen_label = None
for p in doc.paragraphs:
    if p.text.replace(' ', '') == '正文' and len(p.text.strip()) <= 6:
        # 确认是15pt的标签
        for r in p.runs:
            if r.font.size and r.font.size.pt and abs(r.font.size.pt - 15.0) < 1.0:
                zhengwen_label = p; break
        if zhengwen_label is not None:
            break
if zhengwen_label is None:
    # 兜底：关键词后第一个含'正文'的段
    seen_kw = False
    for p in doc.paragraphs:
        if p.text.strip().startswith('【关键词】'):
            seen_kw = True; continue
        if seen_kw and '正' in p.text and '文' in p.text and len(p.text.strip()) <= 6:
            zhengwen_label = p; break
print('正 文 标签:', zhengwen_label.text if zhengwen_label else None)

# 给正文标签加页前分页
if zhengwen_label is not None:
    pPr = zhengwen_label._element.get_or_add_pPr()
    if pPr.find(qn('w:pageBreakBefore')) is None:
        pb = OxmlElement('w:pageBreakBefore'); pPr.append(pb)

# 正文内容块
# ('h1'/'h2'/'body'/'fig'/'table'/'code')
blocks = [
    ('h1','1 绪论'),
    ('h2','1.1 课题背景与研究意义'),
    ('body','随着我国高等教育事业的快速发展，高校招生规模不断扩大，学生宿舍的数量和入住人数持续增加。宿舍作为学生学习和生活的重要场所，其水电、照明、门窗、家具、网络等设施设备在使用过程中不可避免地会出现老化或损坏。宿舍设施能否得到及时维修，直接关系到学生的日常生活体验和学校的后勤服务保障水平。'),
    ('body','然而，目前多数高校的宿舍报修仍以传统方式为主：学生需要到宿舍楼值班室填写纸质报修单，或者通过电话向宿管人员反映，再由宿管人员汇总后通知维修人员上门处理。这种报修方式存在诸多弊端：一是报修流程繁琐，学生填写纸质单据费时费力，信息也容易遗漏或丢失；二是信息传递不及时，从学生报修到维修人员上门往往需要较长时间，缺乏有效的时间约束；三是维修进度无法跟踪，学生难以了解报修处理的实时状态，只能反复询问；四是缺少统一的记录和统计，后勤部门难以对报修数据进行汇总分析，不利于管理决策。'),
    ('body','随着移动互联网技术和智慧校园建设的推进，利用信息化手段改善后勤服务已成为高校发展的必然趋势。微信小程序作为微信生态的重要组成部分，具有免下载安装、即用即走、跨平台运行、开发成本低等优势，非常适合校园场景下的便民服务应用。本文设计并实现的宿舍便民报修微信小程序，通过将报修业务线上化、流程化和闭环化，能够有效解决传统报修方式存在的问题，提高维修响应速度和服务质量，对推进智慧校园建设、提升学生满意度具有重要的现实意义。'),
    ('h2','1.2 国内外研究现状'),
    ('body','近年来，国内不少学者和高校围绕报修系统的信息化建设开展了研究。陈淼等针对传统报修方式容易出现的漏报、错报及接收不及时等问题，设计了基于微信小程序的校园报修管理系统，实现了用户在线报修与后台高效接收[1]；李丹等利用微信小程序云开发技术设计了网络报修系统，通过云函数、云数据库和云存储作为后台服务，实现了报修申请、报修派单、报修历史查看和消息通知等功能[3]；邵嘉诚等设计了基于微信小程序的后勤报修系统，充分利用了手机端拍照、录音和定位等功能，解决了跨平台问题[4]；李艳梅等基于Flask框架和微信小程序技术实现了高校智慧报修系统，有效简化了报修流程[2]。上述研究为宿舍报修系统的建设提供了有益的参考。'),
    ('body','国外高校在设施维修管理方面多采用信息化工单管理系统，例如有研究基于Spring Cloud微服务架构实现了宿舍维修管理系统，涵盖了宿舍信息、维修信息和用户信息的管理功能[10]，但此类系统多为Web端应用，面向学生移动端、紧密结合微信生态的报修应用相对较少。'),
    ('body','综合分析国内外研究现状可以发现，报修系统正朝着移动化、智能化、轻量化方向发展。相比传统前后端分离架构需要独立部署服务器、运维成本较高的问题，微信小程序云开发模式无需自建服务器，能够显著降低开发与运维成本，且天然契合微信庞大的用户生态。因此，本文选择基于微信小程序云开发技术设计宿舍便民报修系统，具有较好的可行性和应用前景。'),
    ('h2','1.3 主要研究内容与论文结构'),
    ('body','本文的主要研究内容包括：（1）对宿舍报修业务进行需求分析，明确系统的功能需求和非功能需求；（2）完成系统的总体架构、功能模块和数据库设计；（3）基于微信小程序云开发技术实现系统各功能模块；（4）对系统进行功能测试，验证系统功能完整性和运行稳定性。'),
    ('body','本文共分为七章，各章内容安排如下：第一章为绪论，介绍课题背景、研究意义、国内外研究现状及主要研究内容；第二章介绍系统开发所涉及的关键技术；第三章对系统进行需求分析；第四章完成系统的总体设计；第五章介绍系统的具体实现；第六章对系统进行测试；第七章对全文进行总结与展望。'),
    ('h1','2 相关技术介绍'),
    ('h2','2.1 微信小程序技术'),
    ('body','微信小程序是腾讯公司于2017年正式推出的一种不需要下载和安装即可使用的应用形态，用户通过微信扫一扫或搜索即可打开使用，体现了"触手可及、用完即走"的产品理念[8]。小程序依托微信庞大的用户生态，天然具备用户基础广、推广成本低的优势，已被广泛应用于电商、生活服务、校园服务等众多领域。'),
    ('body','从技术架构上看，微信小程序采用视图层与逻辑层分离的双线程架构。视图层由WXML和WXSS组成，WXML用于描述页面结构，WXSS用于描述页面样式，负责页面的渲染展示；逻辑层由JavaScript编写，负责业务逻辑处理和与微信原生能力的交互。一个小程序项目通常包含app.js、app.json、app.wxss三个全局文件和若干页面文件，每个页面由.wxml、.wxss、.js、.json四种文件组成[9]。开发者通过微信官方提供的组件和API，可以快速构建功能丰富的应用界面。'),
    ('body','相比传统App，微信小程序具有免安装、跨平台、开发门槛低、更新无需用户手动升级等优点，尤其适合宿舍报修这类使用频率不算很高、但对便捷性要求较高的应用场景。'),
    ('h2','2.2 微信小程序云开发技术'),
    ('body','微信小程序云开发是腾讯云与微信团队联合推出的Serverless云服务能力，它打通了微信生态，为开发者提供云数据库、云存储、云函数、云托管等一站式后端服务[6]。使用云开发后，开发者无需购买和搭建服务器，即可快速实现小程序的业务后端，从而将主要精力放在前端业务逻辑上。'),
    ('body','云数据库是云开发提供的一种文档型数据库，数据以JSON格式存储，采用集合和文档的组织形式，支持索引、事务和丰富的查询能力。开发者无需编写SQL和进行数据库运维，可通过云开发提供的SDK在客户端或云函数中便捷地读写数据。云函数是运行在云端的Node.js代码，用于封装服务端业务逻辑，如身份鉴权、数据校验、复杂查询等，通过wx-server-sdk访问云数据库和云存储，能够保证数据操作的安全性和权限可控性。云存储则用于存放图片、文件等非结构化数据，返回唯一的fileID供客户端引用，适合存放报修现场图片等素材。'),
    ('body','总体而言，云开发具有免运维、按量计费、安全可靠、开发效率高等优点。对于宿舍便民报修这类数据量适中、并发量不大的校园应用，云开发模式能够以较低成本实现完整的后端能力，是本文选择该技术方案的主要原因。贾志杰等在《微信小程序与云开发实践》中系统介绍了小程序与云开发的整合开发方法，为本文的技术实现提供了重要参考[7]。'),
    ('h2','2.3 其他相关技术'),
    ('body','除微信小程序和云开发外，系统开发还涉及以下技术：微信开发者工具是官方提供的一体化集成开发环境，支持代码编辑、实时预览、真机调试、性能分析和代码上传等功能，是系统开发与调试的主要工具；订阅消息是微信提供的消息推送能力，系统通过订阅消息向用户发送报修进度通知，提升用户体验；JavaScript语言及ES6语法、Promise异步编程用于实现页面逻辑和云函数业务处理。'),
    ('h1','3 系统需求分析'),
    ('h2','3.1 系统可行性分析'),
    ('body','从技术角度看，微信小程序开发技术和云开发技术已经非常成熟，微信开发者工具提供了完善的开发、调试和发布流程，云开发提供的云函数、云数据库、云存储能够满足系统对后端服务的全部需求，开发难度适中，技术方案切实可行。'),
    ('body','从经济角度看，云开发采用按量计费模式并提供了免费额度，对于宿舍报修这类中小规模应用，开发和运行成本很低，无需购置独立服务器，也无需配备专职运维人员，经济上完全可行。'),
    ('body','从操作角度看，系统基于微信小程序实现，学生只需在微信中搜索或扫码即可使用，无需额外安装和学习成本；管理端界面简洁直观，管理员和维修人员经过简单培训即可上手操作，操作上可行。'),
    ('body','从管理角度看，系统的应用契合高校后勤信息化建设的发展方向，能够帮助后勤部门实现报修业务的数字化管理，提升服务质量，具备良好的管理环境和支持条件。'),
    ('h2','3.2 系统功能需求分析'),
    ('body','通过对宿舍报修业务的深入调研，系统面向学生用户、维修人员和后勤管理员三类角色，其功能需求分析如下。'),
    ('body','（1）学生用户功能：学生通过微信授权登录系统后，可以查看个人身份信息；在线提交报修申请时，可以选择报修类型（如水电、家具、网络、门窗等）、填写故障描述、上传故障现场图片、选择宿舍楼栋和房间号；提交后可以随时查看报修记录和维修进度，对长时间未处理的工单可以进行催办；维修完成后可以确认完成并对维修服务质量进行评分和评价。'),
    ('body','（2）维修人员功能：维修人员登录后可以查看被指派的报修工单，查看工单的故障详情、报修位置和报修图片；接单后可以上门维修，维修过程中可以更新工单状态，维修完成后填写维修结果和完成说明。'),
    ('body','（3）管理员功能：管理员可以对报修工单进行审核和派单，将工单指派给合适的维修人员，也可以进行改派和关闭处理；可以管理报修类型和维修人员信息；可以对报修数据进行统计分析，查看不同类型、不同状态的报修数量，为后勤管理决策提供数据支持。'),
    ('h2','3.3 系统非功能需求分析'),
    ('body','除功能需求外，系统还需要满足以下非功能需求：'),
    ('body','（1）性能需求：系统应保证较快的响应速度，页面加载和业务操作响应时间应控制在合理范围内，能够支持一定数量的用户同时在线使用，在高峰期不出现明显卡顿。'),
    ('body','（2）安全需求：系统应具备完善的用户身份鉴权机制，通过微信openid识别用户唯一身份，不同角色拥有不同的操作权限；数据传输应采用加密通道，防止用户信息泄露。'),
    ('body','（3）易用性需求：系统界面应简洁友好、操作流程清晰，学生能够快速完成报修操作，管理员能够方便地管理工单。'),
    ('body','（4）可维护性和可扩展性需求：系统代码结构清晰、模块化程度高，便于后期的维护和功能扩展；数据库设计应具备良好的扩展性，能够适应业务规模的增长。'),
    ('h1','4 系统设计'),
    ('h2','4.1 系统总体架构设计'),
    ('body','本系统采用微信小程序云开发架构，整体分为表现层、服务层和基础设施层三个层次，如图1所示。'),
    ('fig','fig1_architecture.png','图1 系统总体架构图'),
    ('body','表现层即微信小程序客户端，面向学生用户和管理人员提供人机交互界面，包括报修提交、报修记录、工单管理、数据统计等页面；服务层基于云开发平台，由云函数、云数据库、云存储和订阅消息等组成，负责业务逻辑处理、数据存储、文件存储和消息通知；基础设施层依托微信生态，提供微信登录、OpenID鉴权、消息推送等底层能力。'),
    ('body','系统的数据流向为：用户在小程序端发起报修请求，客户端通过微信SDK调用云函数，云函数对请求进行身份校验和参数校验后，将报修数据写入云数据库，同时将报修图片上传到云存储；维修处理过程中，云函数根据工单状态变化通过订阅消息向用户推送进度通知；用户和管理员通过小程序端查询和操作相关数据。'),
    ('h2','4.2 系统功能模块设计'),
    ('body','根据需求分析的结果，系统功能划分为学生端和管理端两大模块，各模块的功能结构如图2所示。'),
    ('fig','fig2_module.png','图2 系统功能模块图'),
    ('body','学生端模块包括微信登录授权、在线报修提交、报修记录与进度查询、报修催办、维修确认与评价反馈等功能。其中，在线报修提交是系统的核心功能，用户需要完整填写报修类型、故障描述、宿舍位置等信息，并可选上传故障图片。'),
    ('body','管理端模块包括报修工单管理、维修人员管理、报修类型管理和报修数据统计等功能。工单管理是管理端的核心功能，管理员对工单进行审核、派单、改派和关闭等操作，实现对报修流程的闭环管控。'),
    ('h2','4.3 数据库设计'),
    ('body','系统采用云开发的文档型数据库存储业务数据，根据业务需求设计了用户、报修类型、报修工单、维修人员和评价五个数据集合，各集合之间的主要关系如图3所示。'),
    ('fig','fig4_er.png','图3 数据库主要集合关系图'),
    ('body','各数据集合的主要字段设计如表1所示。'),
    ('table1',),
    ('body','其中，一个用户可以提交多条报修工单，用户与报修工单是一对多关系；一种报修类型可以对应多条报修工单，报修类型与报修工单是一对多关系；一个维修人员可以处理多条报修工单，维修人员与报修工单是一对多关系；一条报修工单对应一条评价，报修工单与评价是一对一关系。'),
    ('body','报修工单的状态在系统内流转，主要包括"待受理""已派单""维修中""待确认""已完成""已关闭"等状态。工单创建后初始为"待受理"状态，管理员审核派单后变为"已派单"，维修人员接单后变为"维修中"，维修完成通知用户后为"待确认"，用户确认完成后变为"已完成"，工单归档或取消时为"已关闭"，报修业务流程如图4所示。'),
    ('fig','fig3_flow.png','图4 报修业务流程图'),
    ('h1','5 系统实现'),
    ('h2','5.1 开发环境与工具'),
    ('body','系统的开发环境与工具主要包括：操作系统为Windows 10及以上；开发工具为微信开发者工具（稳定版），用于小程序的编写、调试与预览；开发语言为JavaScript（ES6）；后端基于微信小程序云开发平台，使用云函数、云数据库和云存储；数据库为云开发文档型数据库；测试环境包括微信开发者工具模拟器和手机微信真机。'),
    ('h2','5.2 学生端主要功能实现'),
    ('body','（1）微信登录功能。学生首次进入小程序时，调用wx.login获取临时登录凭证，再通过云函数获取用户唯一标识openid，并查询或创建用户记录。登录云函数的核心代码如下：'),
    ('code','// 云函数 login\nconst cloud = require(\'wx-server-sdk\')\ncloud.init({ env: cloud.DYNAMIC_CURRENT_ENV })\nconst db = cloud.database()\nexports.main = async (event) => {\n  const { OPENID } = cloud.getWXContext()\n  const users = db.collection(\'user\')\n  const res = await users.where({ openid: OPENID }).get()\n  if (res.data.length === 0) {\n    await users.add({ data: { openid: OPENID, role: \'student\', createTime: db.serverDate() } })\n  }\n  return { openid: OPENID }\n}'),
    ('body','（2）在线报修功能。学生在报修页面选择报修类型、填写故障描述、选择宿舍楼栋和房间号，并可通过chooseMedia接口选择故障现场图片，图片通过云存储上传后返回fileID，最后调用云函数submitRepair将报修数据写入报修工单集合。报修提交云函数的核心代码如下：'),
    ('code','// 云函数 submitRepair\nconst cloud = require(\'wx-server-sdk\')\ncloud.init({ env: cloud.DYNAMIC_CURRENT_ENV })\nconst db = cloud.database()\nexports.main = async (event) => {\n  const { OPENID } = cloud.getWXContext()\n  const { typeId, title, desc, images, building, room } = event\n  const orderNo = \'BX\' + Date.now()  // 生成工单编号\n  return db.collection(\'repair_order\').add({\n    data: {\n      orderNo, userId: OPENID, typeId, title, desc,\n      images: images || [], building, room,\n      status: \'pending\', createTime: db.serverDate()\n    }\n  })\n}'),
    ('body','（3）报修记录与进度查询。学生在"我的报修"页面查询报修工单集合，展示报修记录列表和每条工单的当前状态，并根据状态显示不同的提示信息，方便学生实时了解维修进度。'),
    ('body','（4）维修确认与评价。维修完成后，学生在工单详情页点击"确认完成"，并填写评分和评价内容，系统将评价数据保存到评价集合，同时更新工单状态为"已完成"。'),
    ('h2','5.3 管理端主要功能实现'),
    ('body','（1）工单管理功能。管理员进入工单管理页面后，可以查看所有报修工单，并根据工单状态进行筛选。对"待受理"状态的工单，管理员审核后选择维修人员进行派单，系统更新工单状态为"已派单"并记录维修人员信息。派单云函数的核心代码如下：'),
    ('code','// 云函数 assignOrder\nconst cloud = require(\'wx-server-sdk\')\ncloud.init({ env: cloud.DYNAMIC_CURRENT_ENV })\nconst db = cloud.database()\nexports.main = async (event) => {\n  const { orderId, workerId } = event\n  return db.collection(\'repair_order\').doc(orderId).update({\n    data: { workerId, status: \'assigned\', assignTime: db.serverDate() }\n  })\n}'),
    ('body','（2）维修人员管理功能。管理员可以在管理端对维修人员的姓名、联系电话、技能工种和在职状态等信息进行新增、修改和删除操作，便于派单时选择合适的维修人员。'),
    ('body','（3）报修数据统计功能。系统通过云函数对报修工单集合进行聚合统计，按报修类型和工单状态统计报修数量，并以图形化方式在管理端展示，为后勤管理决策提供数据参考。'),
    ('h1','6 系统测试'),
    ('h2','6.1 测试环境与方法'),
    ('body','系统测试在Windows 10操作系统上进行，前端测试环境为微信开发者工具模拟器和Android手机微信真机，后端为云开发环境。测试方法采用黑盒测试方法，即不关注系统内部实现细节，仅根据功能需求设计测试用例，验证系统各功能模块的输入输出是否符合预期。'),
    ('h2','6.2 功能测试用例与结果'),
    ('body','根据系统功能需求，设计了登录授权、报修提交、图片上传、报修记录查询、工单派单、维修状态更新、评价反馈等主要测试用例，测试用例及结果如表2所示。'),
    ('table2',),
    ('h2','6.3 测试结论'),
    ('body','测试结果表明，系统各功能模块均能按照预期正常工作，报修流程完整闭环，页面响应及时，数据读写准确，未发现影响系统正常使用的功能缺陷。系统达到了设计要求，能够满足宿舍报修业务的日常使用需要。'),
    ('h1','7 总结与展望'),
    ('h2','7.1 工作总结'),
    ('body','本文针对高校宿舍报修流程繁琐、效率低下的问题，设计并实现了一个基于微信小程序云开发的宿舍便民报修系统。主要工作包括：通过对宿舍报修业务的分析，明确了系统在功能和非功能方面的需求；设计了系统的总体架构、功能模块和数据库结构；基于微信小程序云开发技术完成了学生端和管理端各功能模块的开发与实现；对系统进行了功能测试，验证了系统的功能完整性和运行稳定性。'),
    ('body','系统采用微信小程序云开发模式，无需独立部署和维护服务器，降低了开发与运维成本。系统实现了报修业务的线上化、流程化和闭环化管理，有效简化了报修流程，提高了维修响应速度，改善了学生的报修体验，对提升高校后勤服务质量具有积极意义。'),
    ('h2','7.2 不足与展望'),
    ('body','由于时间和能力有限，系统仍存在一些不足之处：一是系统未与学校统一的身份认证系统对接，用户身份信息主要依赖微信授权，可能存在身份识别不完整的问题；二是报修数据统计分析功能较为基础，缺乏更深入的数据挖掘和分析；三是报修图片尚未进行内容审核，可能存在违规图片上传的风险。'),
    ('body','针对上述不足，后续可以从以下几个方面进行改进：一是对接学校统一身份认证平台，实现实名认证和更完善的权限管理；二是利用大数据和机器学习技术，对报修数据进行深度分析，实现维修需求的智能预测和维修资源的智能调度；三是引入图片内容审核机制，加强平台内容安全管理；四是进一步优化界面交互体验，增加报修评价的多样性，不断提升系统的实用性和用户体验。'),
]

# 生成并插入正文元素
def build_table1():
    tbl = doc.add_table(rows=6, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    data = [
        ('数据集合', '主要字段', '说明'),
        ('user（用户）', '_id、openid、昵称、角色、宿舍楼栋、房间号、联系电话', '存储用户基本信息与角色'),
        ('repair_type（报修类型）', '_id、类型名称、图标、排序', '存储报修故障分类'),
        ('repair_order（报修工单）', '_id、工单编号、报修用户、报修类型、故障描述、报修图片、宿舍信息、维修人员、工单状态、创建/更新时间', '存储报修业务数据'),
        ('repair_worker（维修人员）', '_id、姓名、联系电话、技能工种、在职状态', '存储维修人员信息'),
        ('evaluation（评价）', '_id、工单ID、评分、评价内容、创建时间', '存储用户评价数据'),
    ]
    for r_i, row in enumerate(data):
        for c_i, val in enumerate(row):
            cell = tbl.cell(r_i, c_i)
            cell.text = ''
            p = cell.paragraphs[0]
            r = p.add_run(val)
            set_run_font(r, eastasia='宋体', size_pt=10.5)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i != 1 else WD_ALIGN_PARAGRAPH.LEFT
            # 表头底色
            if r_i == 0:
                shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'),'D9D9D9')
                cell._tc.get_or_add_tcPr().append(shd)
                r.font.bold = True
    # 边框
    set_table_borders(tbl)
    return tbl._tbl

def build_table2():
    tbl = doc.add_table(rows=10, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    data = [
        ('测试编号','测试内容','预期结果','实际结果'),
        ('TC-01','学生微信授权登录','获取openid并创建用户记录，登录成功','通过'),
        ('TC-02','学生提交报修（含图片）','报修工单创建成功，图片正常上传','通过'),
        ('TC-03','报修必填项校验','未填关键项时给出提示，无法提交','通过'),
        ('TC-04','学生查看报修记录与进度','正确显示报修记录和工单状态','通过'),
        ('TC-05','报修催办','工单产生催办记录并通知管理员','通过'),
        ('TC-06','管理员审核并派单','工单状态变为已派单并记录维修人员','通过'),
        ('TC-07','维修人员更新维修状态','工单状态按流程正确流转','通过'),
        ('TC-08','学生确认完成并评价','评价数据保存，工单状态变为已完成','通过'),
        ('TC-09','报修数据统计','按类型和状态统计数量正确','通过'),
    ]
    for r_i, row in enumerate(data):
        for c_i, val in enumerate(row):
            cell = tbl.cell(r_i, c_i)
            cell.text = ''
            p = cell.paragraphs[0]
            r = p.add_run(val)
            set_run_font(r, eastasia='宋体', size_pt=10.5)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if r_i == 0:
                shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'),'D9D9D9')
                cell._tc.get_or_add_tcPr().append(shd)
                r.font.bold = True
    set_table_borders(tbl)
    return tbl._tbl

def set_table_borders(tbl):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'000000')
        borders.append(e)
    tblPr.append(borders)

# 组装正文元素
body_elems = []
t1 = None; t2 = None
for blk in blocks:
    kind = blk[0]
    if kind == 'h1':
        body_elems.append(h1_para(blk[1]))
    elif kind == 'h2':
        body_elems.append(h2_para(blk[1]))
    elif kind == 'body':
        body_elems.append(body_para(blk[1]))
    elif kind == 'fig':
        fname, cap = blk[1], blk[2]
        body_elems.append(image_para(fname))
        body_elems.append(caption_para(cap))
    elif kind == 'code':
        body_elems.append(code_para(blk[1]))
    elif kind == 'table1':
        t1 = build_table1()
        body_elems.append(caption_para('表1 系统主要数据表设计'))
        body_elems.append(t1)
    elif kind == 'table2':
        t2 = build_table2()
        body_elems.append(caption_para('表2 主要功能测试用例及结果'))
        body_elems.append(t2)

if zhengwen_label is not None:
    insert_after(zhengwen_label._element, body_elems)

print('阶段2完成：正文内容已插入')

# ================= 8. 参考文献 =================
refs = [
    '[1]陈淼,羊开云.基于微信小程序的校园报修管理系统设计[J].网络安全技术与应用,2024(1):57-60.',
    '[2]李艳梅,董保栋,荣静,等.基于微信小程序的高校智慧报修系统设计与实现[J].电脑知识与技术,2026(3):52-54.',
    '[3]李丹,张未未.基于小程序与云开发的网络报修系统的设计与实现[J].微型电脑应用,2022(12):187-189.',
    '[4]邵嘉诚,厉旭杰,林选,等.基于微信小程序的后勤报修系统的设计与实现[J].智能计算机与应用,2019,9(2):100-103.',
    '[5]李广用,陈伟康.基于微信小程序的网络故障报修系统的设计[J].现代信息科技,2021,5(3):25-27.',
    '[6]丁晓进,朱晓荣.基于微信小程序的云计算学习平台开发[J].软件导刊,2023,22(9).',
    '[7]贾志杰,罗丽红.微信小程序与云开发实践[M].北京:清华大学出版社,2026.',
    '[8]周文洁.微信小程序开发零基础入门[M].北京:清华大学出版社,2019.',
    '[9]李莉,朱壮华,魏秀安.微信小程序开发项目教程(慕课版)(第2版)[M].北京:人民邮电出版社,2025.',
    '[10]SUN J. Design and implementation of dormitory repair management system based on Spring Cloud Microservices[J]. International Core Journal of Engineering, 2021, 7(1):201-204.',
]

# 删除示例参考文献（含续行）
for p in list(doc.paragraphs):
    t = p.text.strip()
    if t.startswith('[1]') or t.startswith('[2]') or t.startswith('[3]') or t.startswith('[4]') or t.startswith('[5]') \
       or '绿色金融' in t or 'Liu Qingrui' in t or 'Finance,2022' in t or 'transformation of energy consumption' in t:
        p._element.getparent().remove(p._element)

# 找到参考文献标题，在其后插入
ref_heading = None
for p in doc.paragraphs:
    if p.text.strip() == '参考文献':
        # 判断是正文末尾的参考文献标题（非目录中）
        ref_heading = p; break
if ref_heading is not None:
    # 清空标题 runs，重新设置为四号宋体加粗居中
    for r in list(ref_heading.runs):
        r._element.getparent().remove(r._element)
    nr = ref_heading.add_run('参考文献')
    set_run_font(nr, eastasia='宋体', size_pt=14, bold=True)
    set_para(ref_heading, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=10)
    # 插入文献
    ref_elems = []
    for ref in refs:
        p_el = OxmlElement('w:p')
        p = Paragraph(p_el, doc)
        r = p.add_run(ref)
        set_run_font(r, eastasia='宋体', ascii_f='Times New Roman', size_pt=10.5)
        set_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_exact=18,
                 left_indent_twips=420, hanging_twips=420)
        ref_elems.append(p_el)
    insert_after(ref_heading._element, ref_elems)

print('阶段3完成：参考文献')

# ================= 9. 评阅表 =================
# 找到评阅表（含"论文性质"与"毕业设计"行的表）
tbl_review = None
for t in doc.tables:
    flat = ' '.join(c.text for row in t.rows for c in row.cells)
    if '论文性质' in flat and '毕业设计' in flat:
        tbl_review = t
        break
if tbl_review is None:
    tbl_review = doc.tables[-1]
print('评阅表定位:', [c.text[:8] for c in tbl_review.rows[0].cells])

def uniq_row_cells(row):
    uniq = []
    seen = set()
    for c in row.cells:
        if id(c._tc) in seen: continue
        seen.add(id(c._tc)); uniq.append(c)
    return uniq

def fill_cell(cell, text, size=12, align=WD_ALIGN_PARAGRAPH.CENTER):
    for pr in list(cell.paragraphs):
        pr._element.getparent().remove(pr._element)
    p = cell.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, eastasia='宋体', size_pt=size)
    p.alignment = align

for r_i, row in enumerate(tbl_review.rows):
    cells = uniq_row_cells(row)
    texts = [c.text.strip() for c in cells]
    joined = '|'.join(texts)
    if '学    院' in texts or '学院' in texts[0]:
        # 学院 与 专业班级 行
        if '专业班级' in texts:
            for ci, c in enumerate(cells):
                if c.text.strip() == '学院' or c.text.strip() == '学    院':
                    if ci+1 < len(cells): fill_cell(cells[ci+1], '电子信息学院')
                if c.text.strip() == '专业班级':
                    if ci+1 < len(cells): fill_cell(cells[ci+1], '软件技术C240313班')
    elif '题' in texts[0] and '目' in texts[0]:
        # 题目行
        for ci, c in enumerate(cells):
            if '题目' in c.text or '题' in c.text:
                if ci+1 < len(cells): fill_cell(cells[ci+1], '宿舍便民报修微信小程序的设计与实现', size=12)
                break
    elif '论文性质' in texts[0] or '论文性质' in joined:
        # 勾选毕业设计
        for c in cells:
            if '毕业设计' in c.text:
                newr = OxmlElement('w:r')
                rPr = OxmlElement('w:rPr')
                rf = OxmlElement('w:rFonts'); rf.set(qn('w:eastAsia'),'宋体'); rPr.append(rf)
                sz = OxmlElement('w:sz'); sz.set(qn('w:val'),'24'); rPr.append(sz)
                newr.append(rPr)
                t = OxmlElement('w:t'); t.text = '√'; newr.append(t)
                p0 = c.paragraphs[0]
                p0._element.insert(0, newr)
                break
    elif '一、主要研究内容' in texts[0] or '一、主要研究内容' in joined:
        content = ('围绕宿舍报修业务流程，设计并实现一个基于微信小程序的宿舍便民报修系统。'
                   '系统采用微信小程序云开发技术，面向学生用户和后勤管理人员，实现在线报修、故障类型分类、'
                   '报修进度查询、催办、评价反馈以及报修工单管理、维修人员派单、维修状态更新等功能。'
                   '主要工作包括系统需求分析、总体架构与数据库设计、学生端与管理端功能实现，以及系统功能测试。')
        target = cells[0]
        for pr in list(target.paragraphs):
            pr._element.getparent().remove(pr._element)
        p = target.add_paragraph()
        r = p.add_run(content)
        set_run_font(r, eastasia='宋体', size_pt=12)
        set_para(p, line_exact=22, first_line_chars=2)

print('阶段4完成：评阅表')

# ================= 10. 页脚页码域 =================
from docx.oxml import parse_xml
from lxml import etree as ET

def make_footer_content():
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    xml = ('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           '<w:pPr><w:jc w:val="center"/><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="宋体"/><w:sz w:val="18"/></w:rPr></w:pPr>'
           '<w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="宋体"/><w:sz w:val="18"/></w:rPr>'
           '<w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
           '<w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="宋体"/><w:sz w:val="18"/></w:rPr>'
           '<w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'
           '<w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="宋体"/><w:sz w:val="18"/></w:rPr>'
           '<w:fldChar w:fldCharType="separate"/></w:r>'
           '<w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="宋体"/><w:sz w:val="18"/></w:rPr>'
           '<w:t>1</w:t></w:r>'
           '<w:r><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="宋体"/><w:sz w:val="18"/></w:rPr>'
           '<w:fldChar w:fldCharType="end"/></w:r>'
           '</w:p>')
    return parse_xml(xml)

# 修改三个 footer part
for sec in doc.sections:
    ftr = sec.footer
    if ftr is not None and not ftr.is_linked_to_previous:
        root = ftr._element
        for child in list(root):
            root.remove(child)
        root.append(make_footer_content())

# ================= 11. 摘要节页码从1开始 =================
# section 3 是摘要节（索引2），设置 pgNumType start=1
secs = doc.sections
if len(secs) >= 3:
    sec3 = secs[2]
    sectPr = sec3._sectPr
    pg = sectPr.find(qn('w:pgNumType'))
    if pg is None:
        pg = OxmlElement('w:pgNumType'); sectPr.append(pg)
    pg.set(qn('w:start'), '1')

print('阶段5完成：页脚页码/分节')

doc.save(TGT)
print('论文已保存:', TGT)

