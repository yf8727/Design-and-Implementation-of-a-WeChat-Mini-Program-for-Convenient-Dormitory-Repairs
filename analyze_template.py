# -*- coding: utf-8 -*-
import docx
from docx import Document

path = r'C:\Users\ZhuanZ（无密码）\Desktop\ffffffffff\毕业论文\专科毕业论文（设计）模板及要求(1).docx'
doc = Document(path)

print('=== 分节信息 ===')
for i, sec in enumerate(doc.sections):
    print(f'节{i}: 起始类型={sec.start_type}, 页宽={sec.page_width}, 页高={sec.page_height}, 上={sec.top_margin} 下={sec.bottom_margin} 左={sec.left_margin} 右={sec.right_margin}')

print()
print('=== 段落中的分页符 ===')
for i, p in enumerate(doc.paragraphs):
    for r in p.runs:
        for child in r._element:
            if child.tag.endswith('}br'):
                br_type = child.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type')
                print(f'  段落[{i}] 有<w:br type={br_type}> 文本: {p.text[:30]}')
    # 检查段前分页
    pPr = p._element.pPr
    if pPr is not None:
        if pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pageBreakBefore') is not None:
            print(f'  段落[{i}] 段前分页: {p.text[:30]}')

print()
print('=== 关键段落run格式 ===')
def get_font_info(p):
    info = []
    for r in p.runs:
        f = r.font
        ea = None
        if r._element.rPr is not None and r._element.rPr.rFonts is not None:
            ea = r._element.rPr.rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia')
        sz = f.size.pt if f.size else None
        info.append({'text': r.text[:15], 'name': f.name, 'ea': ea, 'size': sz, 'bold': f.bold})
    return info

for idx in [7, 14, 29, 36, 42, 51, 56, 59, 63, 74, 89, 101]:
    p = doc.paragraphs[idx]
    print(f'段落[{idx}] "{p.text[:20]}": {get_font_info(p)}')
    pf = p.paragraph_format
    print(f'    对齐={pf.alignment}, 行距规则={pf.line_spacing_rule}, 行距={pf.line_spacing}, 首行缩进={pf.first_line_indent}, 段前={pf.space_before}, 段后={pf.space_after}')

# 检查页脚
print()
print('=== 页脚内容 ===')
for i, sec in enumerate(doc.sections):
    footer = sec.footer
    print(f'节{i} 页脚链接到前节: {footer.is_linked_to_previous}')
    for p in footer.paragraphs:
        print(f'  页脚段落: "{p.text}" 对齐={p.alignment}')
    # 检查页脚是否含PAGE域
    from lxml import etree
    for p in footer.paragraphs:
        if 'PAGE' in p._element.xml:
            print('  包含 PAGE 域: True')

# 检查页眉
print()
print('=== 页眉内容 ===')
for i, sec in enumerate(doc.sections):
    header = sec.header
    print(f'节{i} 页眉: {"|".join(p.text for p in header.paragraphs)}')

# 检查 sections count and 每节的段落归属
print()
print('=== 文档body直接子元素数量 ===')
body = doc.element.body
from lxml import etree
cnt_p = 0
cnt_tbl = 0
cnt_sect = 0
for child in body:
    tag = etree.QName(child).localname
    if tag == 'p':
        cnt_p += 1
    elif tag == 'tbl':
        cnt_tbl += 1
    elif tag == 'sectPr':
        cnt_sect += 1
print(f'段落p={cnt_p}, 表格tbl={cnt_tbl}, sectPr={cnt_sect}')
