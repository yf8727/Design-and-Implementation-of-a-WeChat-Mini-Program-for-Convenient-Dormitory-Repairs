# -*- coding: utf-8 -*-
import docx
from lxml import etree

path = r'C:\Users\ZhuanZ（无密码）\Desktop\ffffffffff\毕业论文\work_tmp\TARGET.docx'
doc = docx.Document(path)
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
print('=== 检查正文一级标题的加粗和字号 ===')
for p in doc.paragraphs:
    t = p.text.strip()
    if t in ('1  绪论', '2  相关技术介绍', '5  系统实现'):
        rpr = p._element.find(W + 'pPr/' + W + 'rPr') 
        bold_marker = '<w:b' in p._element.xml
        sz = None
        # 找第一个run的sz
        for r in p._element.findall(W + 'r'):
            szel = r.find(W + 'rPr/' + W + 'sz')
            if szel is not None:
                sz = szel.get(W + 'val')
                break
        print(f'  "{t}": 含加粗={bold_marker}, 字号={sz} (28=四号14pt)')

print()
print('=== 检查二级标题字号 ===')
for p in doc.paragraphs:
    t = p.text.strip()
    if t in ('1.1  研究背景与意义', '3.2  功能需求分析'):
        sz = None
        for r in p._element.findall(W + 'r'):
            szel = r.find(W + 'rPr/' + W + 'sz')
            if szel is not None:
                sz = szel.get(W + 'val')
                break
        print(f'  "{t}": 字号={sz} (24=小四12pt)')

print()
print('=== 检查正文段落的行距(固定22pt=440) ===')
count_440 = 0
count_body = 0
for p in doc.paragraphs:
    t = p.text.strip()
    if t and len(t) > 40 and not t.startswith('['):
        sp = p._element.find(W + 'pPr/' + W + 'spacing')
        if sp is not None:
            line = sp.get(W + 'line')
            rule = sp.get(W + 'lineRule')
            if line == '440' and rule == 'exact':
                count_440 += 1
            count_body += 1
print(f'  正文长段落总数: {count_body}, 其中固定22pt行距: {count_440}')

print()
print('=== 检查参考文献字号(五号=21)与行距(18pt=360) ===')
for p in doc.paragraphs:
    if p.text.strip().startswith('[1]'):
        sz = None
        for r in p._element.findall(W + 'r'):
            szel = r.find(W + 'rPr/' + W + 'sz')
            if szel is not None:
                sz = szel.get(W + 'val')
                break
        sp = p._element.find(W + 'pPr/' + W + 'spacing')
        line = sp.get(W + 'line') if sp is not None else None
        rule = sp.get(W + 'lineRule') if sp is not None else None
        print(f'  参考文献[1]: 字号={sz}(21=五号10.5pt), 行距={line}(360=18pt) 规则={rule}')
        break
