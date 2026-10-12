# -*- coding: utf-8 -*-
import docx
from lxml import etree

path = r'C:\Users\ZhuanZ（无密码）\Desktop\ffffffffff\毕业论文\专科毕业论文（设计）模板及要求(1).docx'
doc = docx.Document(path)
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

print('=== 包含sectPr的段落 ===')
for i, p in enumerate(doc.paragraphs):
    if p._element.find(W + 'pPr/' + W + 'sectPr') is not None:
        print(f'段落[{i}] 含sectPr 文本="{p.text[:30]}"')

# 也检查表格后、body末尾
body = doc.element.body
for child in body:
    if child.tag == W + 'sectPr':
        print('body末尾有sectPr')

# 找出每个sectPr的页码起始等关键属性
print()
print('=== 每节sectPr详细属性 ===')
sects = []
for i, p in enumerate(doc.paragraphs):
    sectPr = p._element.find(W + 'pPr/' + W + 'sectPr')
    if sectPr is not None:
        sects.append((i, sectPr))
# body末
for child in body:
    if child.tag == W + 'sectPr':
        sects.append(('body-end', child))

for idx, sp in sects:
    pg = sp.find(W + 'pgNumType')
    pgnum = None
    if pg is not None:
        pgnum = pg.get(W + 'start')
    print(f'位置={idx}, 页码起始={pgnum}')
