# -*- coding: utf-8 -*-
"""最终交付文件回读验证"""
import re
from docx import Document

f = r'C:\Users\ZhuanZ（无密码）\Desktop\ffffffffff\毕业论文\宿舍便民报修微信小程序的设计与实现-毕业论文（整合版）.docx'
d = Document(f)
ps = d.paragraphs

print('=== 1. 封面表格 ===')
t = d.tables[0]
for row in t.rows:
    print('  ', [c.text for c in row.cells])

print()
print('=== 2. 评阅表关键信息 ===')
t4 = d.tables[3]
r1 = t4.rows[1]
print('  学生姓名:', r1.cells[1].text, '| 学号:', r1.cells[3].text)
print('  论文性质:', t4.rows[3].cells[1].text[:30])

print()
print('=== 3. 摘要修正 ===')
for p in ps:
    if '三类用户' in p.text:
        print('  OK: 摘要已含"三类用户"')

print()
print('=== 4. 目录SDT条目数 ===')
from lxml import etree
body = d.element.body
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
def qn(tag): return W + tag
def ln(el): return etree.QName(el).localname
def elem_text(e): return ''.join(t.text or '' for t in e.iter(qn('t')))
for ch in body:
    if ln(ch) == 'sdt':
        for sc in ch:
            if ln(sc) == 'sdtContent':
                items = list(sc)
                print('  目录条目数:', len(items))
                # 抽查
                print('  首条:', elem_text(items[0]))
                print('  末条:', elem_text(items[-1]))
        break

print()
print('=== 5. 1.2节研究现状（初稿2）===')
for p in ps:
    if '邵嘉诚等[1]' in p.text:
        print('  OK: 1.2已含初稿2研究现状')

print()
print('=== 6. 5.4节 ===')
for p in ps:
    if p.text.strip() == '5.4 消息通知与进度查询实现':
        print('  OK: 5.4小节存在')

print()
print('=== 7. 3.3节需求编号 ===')
nums = []
for p in ps:
    t = p.text.strip()
    m = re.match(r'^（(\d)）(.+?)需求', t)
    if m and ('性能' in t or '可靠性' in t or '安全' in t or '易用' in t or '可维护' in t):
        nums.append(m.group(1) + '.' + m.group(2))
print('  ', nums)

print()
print('=== 8. 参考文献 ===')
ref_sec = False
cnt = 0
for p in ps:
    t = p.text.strip()
    if t == '参考文献':
        ref_sec = True
        continue
    if ref_sec and t.startswith('['):
        cnt += 1
print('  参考文献条数:', cnt)

print()
print('=== 9. 残留检查 ===')
resid = []
for p in ps:
    t = p.text
    if '请更新目录' in t or '正   文' in t or '此处为一级标题' in t:
        resid.append(t[:30])
print('  残留:', resid if resid else '无')

print()
print('=== 10. 图片数量 ===')
dr = d.element.body.findall('.//' + qn('drawing'))
print('  图片drawing数:', len(dr))
print('  表格数:', len(d.tables))
print('  段落数:', len(ps))
